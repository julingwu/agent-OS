"""Document parsing, code search, and surgical file manipulation tools for AgentOS.

Includes:
- fs_read: line-numbered slicing + automatic delegation to doc_read for binaries/office docs
- fs_edit: surgical search and replace with strict uniqueness check
- glob_search: fast workspace recursive file matching with sensible ignore lists
- grep_search: regex/literal workspace content searching with context lines
- doc_read: rich multi-format document parser (PDF, Word, Excel, PPTX, CSV, HTML, JSON, ZIP/TAR)
"""
from __future__ import annotations

import csv
import fnmatch
import io
import json
import logging
import os
import pathlib
import re
import tarfile
import zipfile
from typing import Any

from . import tools

log = logging.getLogger("app.agent.doc_tools")

IGNORED_DIRS = {
    ".git",
    "node_modules",
    ".venv",
    "venv",
    "__pycache__",
    ".pytest_cache",
    ".cache",
    ".idea",
    ".vscode",
    "dist",
    "build",
    "ui_dist",
}

DOC_EXTENSIONS = {
    ".pdf",
    ".docx",
    ".doc",
    ".xlsx",
    ".xls",
    ".pptx",
    ".ppt",
    ".csv",
    ".zip",
    ".tar",
    ".gz",
}


def _rel_path(target: pathlib.Path) -> str:
    root = tools.WORKSPACE.resolve()
    try:
        return str(target.relative_to(root))
    except Exception:
        return target.name


def _safe_read_text(data: bytes) -> str:
    for enc in ("utf-8", "gb18030", "gbk"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return data.decode("utf-8", errors="replace")


def doc_read(path: str, max_chars: int = 30000, sheet: str | None = None) -> dict[str, Any]:
    """解析各类文档、表格、幻灯片与压缩包内容。支持 PDF, DOCX, XLSX, PPTX, CSV, HTML, JSON, ZIP, TAR。"""
    target = tools._jail(path)
    if not target.is_file():
        return {"ok": False, "error": f"文件不存在或不是普通文件: {path}"}

    limit = max(500, min(int(max_chars or 30000), 200000))
    ext = target.suffix.lower()
    rel = _rel_path(target)

    try:
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(target))
            pages = []
            total_pages = len(reader.pages)
            for idx, p in enumerate(reader.pages):
                text = (p.extract_text() or "").strip()
                pages.append(f"--- 第 {idx + 1}/{total_pages} 页 ---\n{text}")
                if sum(len(x) for x in pages) > limit:
                    pages.append("\n...[截断: 达到字数上限]")
                    break
            content = "\n\n".join(pages)
            return {
                "ok": True,
                "path": rel,
                "doc_type": "pdf",
                "total_pages": total_pages,
                "content": content[:limit],
                "chars": len(content),
            }

        if ext in (".docx", ".doc"):
            import docx
            try:
                from docx.opc.exceptions import PackageNotFoundError
            except ImportError:
                PackageNotFoundError = Exception

            try:
                doc = docx.Document(str(target))
            except PackageNotFoundError:
                return {
                    "ok": False,
                    "path": rel,
                    "error": "检测到老版 Word 97-2003 (.doc) 二进制格式，python-docx 仅支持 .docx。建议将文件另存为 .docx 后解析。",
                }
            except Exception as exc:
                if ext == ".doc" or "package" in str(exc).lower() or "not a word file" in str(exc).lower():
                    return {
                        "ok": False,
                        "path": rel,
                        "error": "检测到老版 Word 97-2003 (.doc) 二进制格式，python-docx 仅支持 .docx。建议将文件另存为 .docx 后解析。",
                    }
                raise

            elements = []
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    elements.append(t)
            for tbl in doc.tables:
                rows_txt = []
                for row in tbl.rows:
                    row_cells = [c.text.strip().replace("\n", " ") for c in row.cells]
                    rows_txt.append("| " + " | ".join(row_cells) + " |")
                if rows_txt:
                    elements.append("\n".join(rows_txt))
            content = "\n\n".join(elements)
            return {
                "ok": True,
                "path": rel,
                "doc_type": "docx",
                "paragraphs": len(doc.paragraphs),
                "tables": len(doc.tables),
                "content": content[:limit],
                "chars": len(content),
            }

        if ext in (".xlsx", ".xls"):
            import openpyxl
            try:
                from openpyxl.utils.exceptions import InvalidFileException
            except ImportError:
                InvalidFileException = Exception

            try:
                wb = openpyxl.load_workbook(str(target), data_only=True)
            except InvalidFileException:
                return {
                    "ok": False,
                    "path": rel,
                    "error": "检测到老版 Excel 97-2003 (.xls) 二进制格式，openpyxl 仅支持 .xlsx。建议将文件另存为 .xlsx 后解析。",
                }
            except Exception as exc:
                if ext == ".xls" or "not a valid xlsx" in str(exc).lower() or "zip" in str(exc).lower() or "badzip" in str(exc).lower():
                    return {
                        "ok": False,
                        "path": rel,
                        "error": "检测到老版 Excel 97-2003 (.xls) 二进制格式，openpyxl 仅支持 .xlsx。建议将文件另存为 .xlsx 后解析。",
                    }
                raise

            sheet_names = wb.sheetnames
            selected_sheet = sheet if (sheet and sheet in sheet_names) else sheet_names[0]
            ws = wb[selected_sheet]

            lines = [f"# 工作表: {selected_sheet} (共 {len(sheet_names)} 个表: {', '.join(sheet_names)})"]
            row_count = 0
            for row in ws.iter_rows(values_only=True):
                if any(c is not None and str(c).strip() != "" for c in row):
                    cell_vals = [str(c) if c is not None else "" for c in row]
                    lines.append("| " + " | ".join(cell_vals) + " |")
                    row_count += 1
                if row_count >= 500 or sum(len(l) for l in lines) > limit:
                    lines.append(f"\n...[截断: 共读取 {row_count} 行]")
                    break
            content = "\n".join(lines)
            return {
                "ok": True,
                "path": rel,
                "doc_type": "xlsx",
                "sheet": selected_sheet,
                "sheets": sheet_names,
                "rows": row_count,
                "content": content[:limit],
                "chars": len(content),
            }

        if ext in (".pptx", ".ppt"):
            # PPTX is a zip archive containing ppt/slides/slide{N}.xml
            slides_text = []
            with zipfile.ZipFile(str(target), "r") as zf:
                slide_names = [n for n in zf.namelist() if re.match(r"ppt/slides/slide\d+\.xml", n)]
                slide_names.sort(key=lambda x: int(re.search(r"\d+", x).group()) if re.search(r"\d+", x) else 0)
                for idx, sname in enumerate(slide_names):
                    xml_bytes = zf.read(sname)
                    # Extract text inside <a:t>...</a:t>
                    matches = re.findall(r"<a:t>(.*?)</a:t>", xml_bytes.decode("utf-8", errors="ignore"), re.DOTALL)
                    clean_text = " ".join([m.strip() for m in matches if m.strip()])
                    slides_text.append(f"--- 幻灯片 {idx + 1}/{len(slide_names)} ---\n{clean_text}")
                    if sum(len(s) for s in slides_text) > limit:
                        slides_text.append("\n...[截断: 达到字数上限]")
                        break
            content = "\n\n".join(slides_text)
            return {
                "ok": True,
                "path": rel,
                "doc_type": "pptx",
                "slides": len(slides_text),
                "content": content[:limit],
                "chars": len(content),
            }

        if ext == ".csv":
            raw_bytes = target.read_bytes()[:limit * 2]
            text = _safe_read_text(raw_bytes)
            rows = []
            reader = csv.reader(io.StringIO(text))
            for idx, r in enumerate(reader):
                if r:
                    cleaned_cells = [c.replace("\r\n", " ").replace("\r", " ").replace("\n", " ").replace("|", "\\|") for c in r]
                    rows.append("| " + " | ".join(cleaned_cells) + " |")
                if idx >= 500:
                    rows.append("...[截断: 已达到 500 行]")
                    break
            content = "\n".join(rows)
            return {
                "ok": True,
                "path": rel,
                "doc_type": "csv",
                "rows": len(rows),
                "content": content[:limit],
                "chars": len(content),
            }

        if ext in (".html", ".htm"):
            from bs4 import BeautifulSoup

            raw_bytes = target.read_bytes()[:limit * 3]
            soup = BeautifulSoup(_safe_read_text(raw_bytes), "html.parser")
            for tag in soup(["script", "style", "svg", "noscript"]):
                tag.decompose()
            title = soup.title.get_text(strip=True) if soup.title else ""
            body_text = soup.get_text(separator="\n", strip=True)
            content = f"# {title}\n\n{body_text}" if title else body_text
            return {
                "ok": True,
                "path": rel,
                "doc_type": "html",
                "title": title,
                "content": content[:limit],
                "chars": len(content),
            }

        if ext == ".json":
            raw_bytes = target.read_bytes()[:limit * 2]
            parsed = json.loads(_safe_read_text(raw_bytes))
            content = json.dumps(parsed, ensure_ascii=False, indent=2)
            return {
                "ok": True,
                "path": rel,
                "doc_type": "json",
                "content": content[:limit],
                "chars": len(content),
            }

        if ext == ".zip":
            with zipfile.ZipFile(str(target), "r") as zf:
                infolist = zf.infolist()
                items = [f"{i.filename} ({i.file_size} bytes)" for i in infolist[:100]]
                if len(infolist) > 100:
                    items.append(f"...[共有 {len(infolist)} 个文件]")
                content = "\n".join(items)
                return {
                    "ok": True,
                    "path": rel,
                    "doc_type": "zip",
                    "file_count": len(infolist),
                    "content": content,
                }

        if ext in (".tar", ".gz", ".tgz"):
            with tarfile.open(str(target), "r:*") as tf:
                members = tf.getmembers()
                items = [f"{m.name} ({m.size} bytes)" for m in members[:100]]
                if len(members) > 100:
                    items.append(f"...[共有 {len(members)} 个文件]")
                content = "\n".join(items)
                return {
                    "ok": True,
                    "path": rel,
                    "doc_type": "tar",
                    "file_count": len(members),
                    "content": content,
                }

        # Fallback to plain text read
        raw_bytes = target.read_bytes()[:limit]
        content = _safe_read_text(raw_bytes)
        return {
            "ok": True,
            "path": rel,
            "doc_type": "text",
            "content": content[:limit],
            "chars": len(content),
        }

    except Exception as exc:
        log.exception("doc_read failed for %s", path)
        return {"ok": False, "path": rel, "error": f"解析文档失败: {exc}"}


def fs_read(path: str, start_line: int | None = None, end_line: int | None = None,
            max_bytes: int = 200_000) -> dict[str, Any]:
    """读取工作区里的文件。

    - 若为 Office / PDF / 压缩包 等非纯文本文件，自动委派给 doc_read 提取正文内容。
    - 支持按行号切片（1-indexed），方便精确定位与查看长文件。
    - 兼容返回原有字段格式。
    """
    target = tools._jail(path)
    if not target.is_file():
        return {"ok": False, "error": "not a file"}

    ext = target.suffix.lower()
    if ext in DOC_EXTENSIONS:
        return doc_read(path, max_chars=int(max_bytes or 200_000) // 2)

    limit_bytes = max(1, min(int(max_bytes or 200_000), 2_000_000))
    raw_data = target.read_bytes()[:limit_bytes]

    try:
        text = raw_data.decode("utf-8")
    except UnicodeDecodeError:
        try:
            text = raw_data.decode("gb18030")
        except (UnicodeDecodeError, LookupError):
            return {"ok": False, "error": "binary file (not utf-8/gb18030)"}

    lines = text.splitlines(keepends=True)
    total_lines = len(lines)
    rel = _rel_path(target)

    # If line slicing is requested (1-indexed)
    if start_line is not None or end_line is not None:
        s = max(1, int(start_line or 1))
        e = min(total_lines, int(end_line or total_lines))
        if s > total_lines:
            sliced_lines = []
        else:
            sliced_lines = lines[s - 1:e]

        # Format with line numbers like "1: content"
        formatted_content = "".join(f"{s + idx}: {l}" for idx, l in enumerate(sliced_lines))
        return {
            "ok": True,
            "path": rel,
            "content": formatted_content,
            "start_line": s,
            "end_line": e,
            "total_lines": total_lines,
        }

    return {
        "ok": True,
        "path": rel,
        "content": text,
        "total_lines": total_lines,
    }


def fs_edit(path: str, old_text: str, new_text: str, replace_all: bool = False) -> dict[str, Any]:
    """精确修改工作区文本文件。

    采用外科手术式（Surgical）字符串替换：
    - old_text 必须在目标文件中精准存在且默认具有全局唯一性。
    - 若出现多处匹配，必须传入 replace_all=True 或补充更长的上下文，防止意外多改。
    """
    target = tools._jail(path)
    if not target.is_file():
        return {"ok": False, "error": f"file not found: {path}"}

    if not old_text:
        return {"ok": False, "error": "old_text cannot be empty"}

    raw_bytes = target.read_bytes()
    try:
        content = raw_bytes.decode("utf-8")
        encoding = "utf-8"
    except UnicodeDecodeError:
        try:
            content = raw_bytes.decode("gb18030")
            encoding = "gb18030"
        except (UnicodeDecodeError, LookupError):
            return {"ok": False, "error": "file is binary or cannot be decoded as utf-8/gb18030"}

    count = content.count(old_text)
    if count == 0:
        return {
            "ok": False,
            "error": "old_text not found in file. Ensure exact match including whitespace and indentation.",
        }

    if count > 1 and not replace_all:
        return {
            "ok": False,
            "error": (
                f"old_text matches {count} occurrences in {path}. "
                "Please provide more surrounding context lines to make it unique, or set replace_all=True."
            ),
        }

    if replace_all:
        new_content = content.replace(old_text, new_text)
        replaced_count = count
    else:
        new_content = content.replace(old_text, new_text, 1)
        replaced_count = 1

    target.write_bytes(new_content.encode(encoding))
    return {
        "ok": True,
        "path": _rel_path(target),
        "replacements": replaced_count,
        "bytes": len(new_content.encode(encoding)),
    }


def glob_search(pattern: str, path: str = ".") -> dict[str, Any]:
    """在工作区内根据通配符模式（如 *.py, **/*.json）快速递归搜索文件。"""
    root = tools.WORKSPACE.resolve()
    base_dir = tools._jail(path)

    if not base_dir.exists():
        return {"ok": False, "error": f"directory not found: {path}"}

    matches: list[dict[str, Any]] = []
    max_results = 200

    for current_dir, dirs, files in os.walk(str(base_dir)):
        # Prune ignored directories
        dirs[:] = [d for d in dirs if d not in IGNORED_DIRS and not d.startswith(".")]

        rel_cur = os.path.relpath(current_dir, str(root))

        for filename in files:
            rel_file = filename if rel_cur == "." else os.path.join(rel_cur, filename)
            # Normalize to forward slashes for uniform matching
            norm_path = rel_file.replace("\\", "/")

            if fnmatch.fnmatch(norm_path, pattern) or fnmatch.fnmatch(filename, pattern):
                full_p = pathlib.Path(current_dir) / filename
                try:
                    size = full_p.stat().st_size
                except Exception:
                    size = 0
                matches.append({
                    "path": norm_path,
                    "size": size,
                    "is_dir": False,
                })
                if len(matches) >= max_results:
                    break
        if len(matches) >= max_results:
            break

    return {
        "ok": True,
        "pattern": pattern,
        "base_path": _rel_path(base_dir),
        "count": len(matches),
        "matches": matches,
    }


def grep_search(pattern: str, path: str = ".", file_glob: str = "*",
                context_lines: int = 1, max_matches: int = 80) -> dict[str, Any]:
    """在工作区文本文件中搜索正则表达式或文本片段，返回行号与上下文。"""
    root = tools.WORKSPACE.resolve()
    base_dir = tools._jail(path)

    if not base_dir.exists():
        return {"ok": False, "error": f"directory not found: {path}"}

    try:
        regex = re.compile(pattern, re.IGNORECASE)
    except re.error as e:
        return {"ok": False, "error": f"Invalid regular expression: {e}"}

    matches: list[dict[str, Any]] = []
    limit = max(1, min(int(max_matches or 80), 300))
    ctx = max(0, min(int(context_lines or 1), 5))

    for current_dir, dirs, files in os.walk(str(base_dir)):
        dirs[:] = [d for d in dirs if d not in IGNORED_DIRS and not d.startswith(".")]
        rel_cur = os.path.relpath(current_dir, str(root))

        for filename in files:
            if not fnmatch.fnmatch(filename, file_glob):
                continue

            full_p = pathlib.Path(current_dir) / filename
            ext = full_p.suffix.lower()
            if ext in DOC_EXTENSIONS and ext not in (".csv", ".json", ".html", ".htm"):
                continue

            try:
                raw = full_p.read_bytes()
                if b"\x00" in raw[:1024]:
                    continue  # binary
                text = _safe_read_text(raw)
                lines = text.splitlines()

                rel_file = (filename if rel_cur == "." else os.path.join(rel_cur, filename)).replace("\\", "/")

                for idx, line in enumerate(lines):
                    if regex.search(line):
                        c_start = max(0, idx - ctx)
                        c_end = min(len(lines), idx + ctx + 1)
                        context = [
                            {"line": i + 1, "text": lines[i], "match": i == idx}
                            for i in range(c_start, c_end)
                        ]
                        matches.append({
                            "path": rel_file,
                            "line": idx + 1,
                            "text": line,
                            "context": context,
                        })
                        if len(matches) >= limit:
                            break
            except Exception:
                continue

            if len(matches) >= limit:
                break
        if len(matches) >= limit:
            break

    return {
        "ok": True,
        "pattern": pattern,
        "count": len(matches),
        "matches": matches,
    }
